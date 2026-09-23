#!/usr/bin/env python3
"""Erzeugt aus public/data/modelle.json wieder eine Excel-Datei.

Aufruf:  python3 scripts/json_to_xlsx.py [ziel.xlsx]
Ergebnis: ein Blatt je Bauart, sortiert nach Leistungszahl, mit Kostenformeln.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WURZEL = Path(__file__).resolve().parent.parent
QUELLE = WURZEL / "public" / "data" / "modelle.json"

SPALTEN = [
    ("Modell", "name", 46, None),
    ("Marke", "marke", 16, None),
    ("Liter", "volumen_l", 9, "#,##0"),
    ("ηwh (%)", "eta_wh", 10, "0.0"),
    ("Leistungszahl", "_lz", 14, "0.00"),
    ("Basis", "_basis", 9, None),
    ("Klasse", "eff_klasse", 9, None),
    ("Preis", "preis_eur", 13, '#,##0.00 "€"'),
    ("Schall dB(A)", "schallleistung_db", 13, "#,##0.0"),
    ("Kältemittel", "kaeltemittel", 13, None),
    ("Wärmetauscher", "_wt", 22, None),
    ("Heizstab (W)", "heizstab_w", 13, "#,##0"),
    ("Kesselmaterial", "kessel_material", 32, None),
    ("Anode", "anode", 32, None),
    ("Quelle", "quelle", 40, None),
]
REIHE = ["eta_wh", "cop_a20", "cop_a15", "cop_a14", "cop_a7", "cop_unspezifisch"]
KURZ = {"eta_wh": "ηwh", "cop_a20": "A20", "cop_a15": "A15", "cop_a14": "A14",
        "cop_a7": "A7", "cop_unspezifisch": "k. A."}


def leistungszahl(m, faktor):
    for feld in REIHE:
        wert = round(m["eta_wh"] / 100 * faktor, 3) if feld == "eta_wh" and m.get("eta_wh") else m.get(feld)
        if wert is not None:
            return wert, KURZ[feld]
    return None, "–"


def wt_text(m):
    wt = m.get("waermetauscher") or {}
    if wt.get("vorhanden") is None:
        return "–"
    if not wt["vorhanden"]:
        return "nein"
    teile = [wt.get("variante"), f"{wt['flaeche_m2']:.2f} m²".replace(".", ",") if wt.get("flaeche_m2") else None]
    teile = [t for t in teile if t]
    return "ja" + (f" ({', '.join(teile)})" if teile else "")


def blatt(wb, titel, modelle, faktor, erstes):
    ws = wb.active if erstes else wb.create_sheet()
    ws.title = titel

    kopf_fuellung = PatternFill("solid", fgColor="1A5F8A")
    ws["A1"] = "Brauchwasserwärmepumpen – " + titel
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Strompreis (€/kWh)"
    ws["B2"] = 0.30
    ws["B2"].number_format = '0.00 "€"'
    ws["C2"] = "Bedarf (kWh/a)"
    ws["D2"] = 2000
    ws["E2"] = "Jahreskosten = Bedarf × Strompreis ÷ Leistungszahl"
    ws["E2"].font = Font(italic=True, size=9)

    for i, (label, _, breite, _) in enumerate(SPALTEN, start=1):
        c = ws.cell(row=4, column=i, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = kopf_fuellung
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = breite
    spalte_jk = len(SPALTEN) + 1
    c = ws.cell(row=4, column=spalte_jk, value="Jahreskosten (€/a)")
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = kopf_fuellung
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[get_column_letter(spalte_jk)].width = 18

    modelle = sorted(modelle, key=lambda m: leistungszahl(m, faktor)[0] or -1, reverse=True)
    for zeile, m in enumerate(modelle, start=5):
        lz, basis = leistungszahl(m, faktor)
        werte = {"_lz": lz, "_basis": basis, "_wt": wt_text(m)}
        for i, (_, feld, _, fmt) in enumerate(SPALTEN, start=1):
            wert = werte.get(feld, m.get(feld))
            c = ws.cell(row=zeile, column=i, value=wert if wert is not None else "–")
            if fmt and isinstance(wert, (int, float)):
                c.number_format = fmt
            c.alignment = Alignment(horizontal="left" if i in (1, 13, 14, 15) else "center",
                                    vertical="center")
        lz_spalte = get_column_letter(SPALTEN.index(("Leistungszahl", "_lz", 14, "0.00")) + 1)
        c = ws.cell(row=zeile, column=spalte_jk,
                    value=f'=IFERROR($D$2*$B$2/{lz_spalte}{zeile},"–")')
        c.number_format = '#,##0 "€"'
        c.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(spalte_jk)}{4 + len(modelle)}"
    return len(modelle)


def main(ziel):
    daten = json.loads(QUELLE.read_text(encoding="utf-8"))
    faktor = daten["scop_faktor"]
    wb = Workbook()
    n1 = blatt(wb, "bodenstehend", [m for m in daten["modelle"] if m["bauart"] == "bodenstehend"], faktor, True)
    n2 = blatt(wb, "wandhängend", [m for m in daten["modelle"] if m["bauart"] == "wandhaengend"], faktor, False)
    n3 = blatt(wb, "ohne Kessel", [m for m in daten["modelle"] if m["bauart"] == "ohne_kessel"], faktor, False)
    wb.save(ziel)
    print(f"{ziel}: {n1} bodenstehende, {n2} wandhängende, {n3} ohne Kessel")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "brauchwasserwaermepumpen.xlsx")
