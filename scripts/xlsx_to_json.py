#!/usr/bin/env python3
"""Konvertiert die Excel-Übersicht in die JSON-Datenquelle der Vergleichsseite.

Wichtigste Aufgabe: Die COP-Werte, die in der Excel-Tabelle in einer einzigen
SCOP-Spalte mit gemischten Messbasen stehen, werden in getrennte Felder
(cop_a7 / cop_a15 / cop_a20 / eta_wh) aufgeteilt. Nur so sind die Geräte
untereinander vergleichbar.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

MARKEN = [
    "OCHSNER", "Ochsner", "Stiebel Eltron", "Viessmann", "Bösch", "Austria Email",
    "Hoval", "KWB", "Junkers Bosch", "Bosch", "Buderus", "Herz", "Vaillant",
    "Saunier Duval", "Ovum", "Dimplex", "Panasonic", "Haier", "Remko", "Daikin",
    "Tesy", "LG", "Alarko", "Wolf", "AEG", "Kermi", "Styleboiler", "Hofman Energy",
    "ÖkoFEN", "Brötje", "Remeha", "Kronoterm", "ELCO", "Hisense", "Weishaupt",
]


def slug(text):
    text = str(text).lower()
    for a, b in [("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss"), ("°", ""), ("²", "2"),
                 ("+", "-plus")]:
        text = text.replace(a, b)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text


def zahl(v):
    """Excel-Zelle -> float oder None ('–' und Text werden zu None)."""
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        m = re.fullmatch(r"\s*([\d.]+),?(\d*)\s*", v.replace(".", ""))
        if m:
            return float(f"{m.group(1)}.{m.group(2) or 0}")
    return None


def de_float(s):
    return float(s.replace(",", "."))


def marke_von(name):
    for m in MARKEN:
        if str(name).startswith(m):
            return m
    return str(name).split()[0]


# --- Zuordnung des Haupt-SCOP-Werts zu seiner Messbasis --------------------
BASIS_REGELN = [
    (r"ηwh_m × Faktor", "eta_wh"),
    (r"COP A7\b|COP A7/W55", "cop_a7"),
    (r"COP A14|COP 14 ?°C|A14/W5[35]", "cop_a14"),
    (r"COP A15|COP 15 ?°C|15/12 ?°C|COP A15/W10-55|COP A15/W55|COP A15/W35", "cop_a15"),
    (r"COP A20|COP L20|Raumluft|COP 20/15 ?°C|A20/W10-5[35]|A20/W35", "cop_a20"),
]


def basis_von(quelle):
    q = str(quelle)
    for muster, feld in BASIS_REGELN:
        if re.search(muster, q):
            return feld
    return "cop_unspezifisch"


# COP-Werte stehen immer mit Dezimalkomma (3,09), ηwh-Angaben als ganze Prozent (149 %).
# Das unterscheidet die beiden Fälle zuverlässig.
ZUSATZ_REGELN = [
    (r"A7:\s*(\d+,\d+)", "cop_a7"),
    (r"A15/W10:\s*(\d+,\d+)", "cop_a15"),
    (r"Außenluft:\s*(\d+,\d+)", "cop_a7"),
    (r"COP lt\. Hersteller (\d+,\d+) bei A20", "cop_a20"),
    (r"COP ([\d,]+) bei L21", "cop_a20"),
    (r"([\d,]+) bei L15", "cop_a15"),
]
ZUSATZ_ETA = [
    (r"A14:\s*(\d+)\s*%", "eta_wh_a14"),
    (r"A7:\s*(\d+)\s*%", "eta_wh_a7"),
]


def lies_blatt(ws, bauart):
    modelle = []
    r = 7
    while ws.cell(row=r, column=1).value:
        name = str(ws.cell(row=r, column=1).value).strip()
        quelle = str(ws.cell(row=r, column=13).value or "")
        eta = zahl(ws.cell(row=r, column=5).value)
        scop_zelle = ws.cell(row=r, column=6).value
        faktor = float(ws["B4"].value)

        eintrag = {
            "id": slug(name),
            "name": name,
            "marke": marke_von(name),
            "bauart": bauart,
            "volumen_l": zahl(ws.cell(row=r, column=4).value),
            "eta_wh": round(eta * 100, 1) if eta else None,
            "eff_klasse": str(ws.cell(row=r, column=3).value or "").strip() or None,
            "schallleistung_db": zahl(ws.cell(row=r, column=7).value),
            "bereitschaftsverlust_kwh_tag": zahl(ws.cell(row=r, column=8).value),
            "kaeltemittel": (str(ws.cell(row=r, column=9).value or "").strip() or None),
            "preis_eur": zahl(ws.cell(row=r, column=10).value),
            "heizstab_w": zahl(ws.cell(row=r, column=12).value),
            "kessel_material": (str(ws.cell(row=r, column=14).value or "").strip() or None),
            "anode": (str(ws.cell(row=r, column=15).value or "").strip() or None),
            "ehpa_guetesiegel": (str(ws.cell(row=r, column=2).value or "").strip() or None),
            "quelle": quelle or None,
            "cop_a7": None, "cop_a15": None, "cop_a14": None, "cop_a20": None,
            "cop_unspezifisch": None,
        }
        for feld in ("kaeltemittel", "kessel_material", "anode", "ehpa_guetesiegel"):
            if eintrag[feld] in ("–", "-", "---", ""):
                eintrag[feld] = None

        # Wärmetauscher
        wt_roh = str(ws.cell(row=r, column=11).value or "")
        flaeche = re.search(r"WT ([\d,]+) m²", quelle)
        variante = re.search(r"ja \(([^)]+)\)", wt_roh)
        eintrag["waermetauscher"] = {
            "vorhanden": wt_roh.startswith("ja"),
            "variante": variante.group(1) if variante else None,
            "flaeche_m2": de_float(flaeche.group(1)) if flaeche else None,
        } if wt_roh not in ("–", "") else {"vorhanden": None, "variante": None, "flaeche_m2": None}

        # Haupt-SCOP der passenden Basis zuordnen
        if isinstance(scop_zelle, str) and scop_zelle.startswith("="):
            eintrag["scop_basis"] = "eta_wh"
        else:
            wert = zahl(scop_zelle)
            feld = basis_von(quelle)
            eintrag["scop_basis"] = feld
            if wert is not None:
                eintrag[feld] = wert

        # zusätzlich genannte COP-/ηwh-Werte aus der Quellenspalte holen
        for muster, feld in ZUSATZ_REGELN:
            m = re.search(muster, quelle)
            if m and eintrag.get(feld) is None:
                eintrag[feld] = de_float(m.group(1))
        for muster, feld in ZUSATZ_ETA:
            m = re.search(muster, quelle)
            if m:
                eintrag[feld] = float(m.group(1))

        # abgeleiteter Vergleichswert wie in der Excel-Tabelle
        eintrag["scop_tabelle"] = round(eta * faktor, 3) if eta else (
            eintrag[eintrag["scop_basis"]] if eintrag["scop_basis"] in eintrag else None)

        modelle.append(eintrag)
        r += 1
    return modelle


def main(xlsx, ziel):
    wb = openpyxl.load_workbook(xlsx)
    modelle = lies_blatt(wb["Übersicht"], "bodenstehend")
    modelle += lies_blatt(wb["Wandgeräte"], "wandhaengend")

    ids = [m["id"] for m in modelle]
    doppelt = {i for i in ids if ids.count(i) > 1}
    if doppelt:
        print("WARNUNG doppelte IDs:", doppelt, file=sys.stderr)

    daten = {
        "stand": "2026-09-23",
        "scop_faktor": float(wb["Übersicht"]["B4"].value),
        "quellen_hinweis": (
            "GET-Produktdatenbank (Gruppe 396), Herstellerdatenblätter und Händlerpreise. "
            "COP-Werte sind nach Messbasis getrennt, weil Hersteller unterschiedliche "
            "Lufteintrittstemperaturen angeben."
        ),
        "modelle": modelle,
    }
    Path(ziel).write_text(json.dumps(daten, ensure_ascii=False, indent=2), encoding="utf-8")
    return daten


if __name__ == "__main__":
    d = main(sys.argv[1], sys.argv[2])
    m = d["modelle"]
    print(f"{len(m)} Modelle geschrieben nach {sys.argv[2]}")
    for feld in ("eta_wh", "cop_a7", "cop_a14", "cop_a15", "cop_a20", "cop_unspezifisch",
                 "preis_eur", "kessel_material", "anode", "heizstab_w"):
        print(f"  {feld:<22} befüllt: {sum(1 for x in m if x.get(feld) is not None):>3}/{len(m)}")
